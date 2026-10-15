"""Exercise workbook preservation, input validation, and legacy compatibility."""
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from batch_datasets import (DEFAULT_BATCH_WORKBOOK, DEFAULT_WORKBOOK, REPLACEMENTS,
                            SHEETS, replace_rover_sheets, sheet_snapshot, sheet_spec)
from GPy_true_field import ROVERS, load_sheet_data, load_training_data, train_model
from run import compare_references


class BatchTests(unittest.TestCase):
    def test_replacement_preserves_original_package_and_donor_data(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'datasets.xlsx'
            shutil.copy2(DEFAULT_BATCH_WORKBOOK, path)
            before = path.read_bytes()
            result = replace_rover_sheets(path)
            self.assertEqual(Path(result['backup']).read_bytes(), before)
            book, donor = load_workbook(path), load_workbook(DEFAULT_WORKBOOK)
            try:
                self.assertEqual(tuple(book.sheetnames), SHEETS)
                for dst, src in REPLACEMENTS.items():
                    self.assertEqual(sheet_snapshot(book[dst]), sheet_snapshot(donor[src]))
            finally:
                book.close()
                donor.close()
            with zipfile.ZipFile(result['backup']) as old, zipfile.ZipFile(path) as new:
                for i, name in enumerate(SHEETS, 1):
                    if name not in REPLACEMENTS:
                        member = f'xl/worksheets/sheet{i}.xml'
                        self.assertEqual(old.read(member), new.read(member))
            self.assertEqual(sheet_spec(path, 'test1')['origin'], (65, 5))
            self.assertEqual(sheet_spec(path, 'train2')['origin'], (60, -5))

    def test_bad_observation_and_non_30_degree_region_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'bad.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = 'train2'
            sheet.append(['longitude(60-90)', 'latitude(-5-25)', 'Br', 'B2', 'B3'])
            sheet.append([65, 0, 1, None, 2])
            book.save(path)
            spec = sheet_spec(path, 'train2')
            with self.assertRaisesRegex(ValueError, '第 2 行'):
                load_sheet_data(path, spec)
            sheet['A1'] = 'longitude(60-100)'
            book.save(path)
            with self.assertRaisesRegex(ValueError, '30°'):
                sheet_spec(path, 'train2')
            book.close()

    def test_original_rover_training_matches_legacy_predictions(self):
        models = {}
        for rover in ROVERS:
            X, Y = load_training_data(DEFAULT_WORKBOOK, rover)
            models[rover] = train_model(X, Y)
        comparison = compare_references(models, DEFAULT_WORKBOOK)
        for result in comparison.values():
            self.assertTrue(result['training_data_matches_legacy_model'])
            # Allow numerical solver variation far below the workbook precision.
            self.assertLess(result['prediction_max_abs_difference'], 1e-5)


if __name__ == '__main__':
    unittest.main()
