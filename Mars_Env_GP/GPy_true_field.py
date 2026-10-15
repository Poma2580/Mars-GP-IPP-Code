"""Train one shared-kernel GPy regression model per Mars rover."""

from __future__ import annotations

import argparse
import json
import pickle
from pathlib import Path

import GPy
import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "data" / "Mars_Rover_Magnetic_Field.xlsx"
DEFAULT_OUTPUT = ROOT / "output"
COMPONENTS = ("Br", "B2", "B3")
ROVERS = {
    "zhurong": {
        "sheet": "Zhurong_Magnetic_Field",
        "longitude": "longitude(95-125)",
        "latitude": "latitude(10-40)",
        "origin": (95.0, 10.0),
        "legacy_model": "model_field.pkl",
    },
    "perseverance": {
        "sheet": "Perseverance_Magnetic_Field",
        "longitude": "longitude(65-95)",
        "latitude": "latitude(5-35)",
        "origin": (65.0, 5.0),
        "legacy_model": "model_field_v2.pkl",
    },
}


def load_training_data(workbook_path: Path, rover: str) -> tuple[np.ndarray, np.ndarray]:
    """Read the original five numeric columns, preserving workbook row order."""
    spec = ROVERS[rover]
    return load_sheet_data(workbook_path, spec)


def load_sheet_data(workbook_path: Path, spec: dict) -> tuple[np.ndarray, np.ndarray]:
    """Read a sheet using explicit coordinate columns and its regional origin."""
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = book[spec["sheet"]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        required = (spec["longitude"], spec["latitude"], *COMPONENTS)
        missing = [name for name in required if name not in header]
        if missing:
            raise ValueError(f"{sheet.title} 缺少列: {', '.join(missing)}")
        indices = [header.index(name) for name in required]
        values = []
        for row_number, row in enumerate(rows, start=2):
            selected = [row[index] if index < len(row) else None for index in indices]
            if all(value is None for value in selected):
                continue
            try:
                numeric = [float(value) for value in selected]
            except (TypeError, ValueError) as error:
                raise ValueError(f"{sheet.title} 第 {row_number} 行含非数值或空值") from error
            if not np.isfinite(numeric).all():
                raise ValueError(f"{sheet.title} 第 {row_number} 行含非有限值")
            values.append(numeric)
        if not values:
            raise ValueError(f"{sheet.title} 没有有效观测值")
    finally:
        book.close()

    data = np.asarray(values, dtype=np.float64)
    origin = np.asarray(spec["origin"], dtype=np.float64)
    return data[:, :2] - origin, data[:, 2:]


def train_model(X: np.ndarray, Y: np.ndarray) -> GPy.models.GPRegression:
    """Keep the kernel and one optimize call from the original GPy script."""
    kernel = GPy.kern.RBF(input_dim=2, variance=1.0, lengthscale=1.0)
    model = GPy.models.GPRegression(X, Y, kernel)
    model.optimize(messages=False)
    return model


def train_all(workbook_path: Path = DEFAULT_WORKBOOK, output_dir: Path = DEFAULT_OUTPUT):
    output_dir.mkdir(parents=True, exist_ok=True)
    models = {}
    parameters = {}
    for rover, spec in ROVERS.items():
        X, Y = load_training_data(workbook_path, rover)
        model = train_model(X, Y)
        with (output_dir / f"{rover}.pkl").open("wb") as target:
            pickle.dump(model, target, protocol=pickle.HIGHEST_PROTOCOL)
        models[rover] = model
        parameters[rover] = {
            "sheet": spec["sheet"],
            "components": list(COMPONENTS),
            "training_points": len(X),
            "origin_degrees": list(spec["origin"]),
            "kernel": "GPy.kern.RBF(input_dim=2, variance=1.0, lengthscale=1.0)",
            "optimized_kernel_variance": float(model.kern.variance.values[0]),
            "optimized_lengthscale_degrees": float(model.kern.lengthscale.values[0]),
            "optimized_noise_variance": float(model.Gaussian_noise.variance.values[0]),
        }
        print(f"{rover}: {len(X)} 个点，模型保存到 {output_dir / (rover + '.pkl')}")
    (output_dir / "kernel_parameters.json").write_text(
        json.dumps(parameters, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return models


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    train_all(args.workbook, args.output_dir)


if __name__ == "__main__":
    main()
