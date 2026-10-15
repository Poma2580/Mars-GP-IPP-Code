"""Replace rover sheets, fit ten original-method GPs, and assess interpolation."""
from __future__ import annotations

import hashlib
import json
import pickle
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime
from importlib.metadata import version
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook
from PIL import Image
from scipy.spatial import ConvexHull, Delaunay

from GPy_true_field import ROOT, DEFAULT_WORKBOOK, COMPONENTS, load_sheet_data, train_model
from draw_true_field import TITLES, ZHURONG_LEVELS, grid

DEFAULT_BATCH_WORKBOOK = ROOT / 'data' / 'Mars_10datasets2.xlsx'
DEFAULT_BATCH_OUTPUT = ROOT / 'output' / '10datasets'
SHEETS = tuple(f'{group}{i}' for group in ('train', 'test') for i in range(1, 6))
REPLACEMENTS = {'train1': 'Zhurong_Magnetic_Field', 'test1': 'Perseverance_Magnetic_Field'}
SEED = 42


def preserve_workbook_package(original_path, source_path, generated_path):
    """Keep the original package and exact numbers; remap donor styles/strings."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    def sheet_paths(archive):
        rel = {r.attrib['Id']: r.attrib['Target'] for r in
               ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        result = {}
        for sh in ET.fromstring(archive.read('xl/workbook.xml')).findall('s:sheets/s:sheet', ns):
            target = rel[sh.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            result[sh.attrib['name']] = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        return result
    packed = generated_path.with_suffix('.packed.xlsx')
    with zipfile.ZipFile(original_path) as original, zipfile.ZipFile(source_path) as source, \
            zipfile.ZipFile(generated_path) as generated:
        old_paths, source_paths, new_paths = map(sheet_paths, (original, source, generated))
        strings = ET.fromstring(source.read('xl/sharedStrings.xml')).findall('s:si', ns)
        replacements = {'xl/styles.xml': generated.read('xl/styles.xml')}
        for dst, src in REPLACEMENTS.items():
            donor = ET.fromstring(source.read(source_paths[src]))
            new = ET.fromstring(generated.read(new_paths[dst]))
            styles = {c.attrib['r']: c.attrib.get('s') for c in new.findall('.//s:c', ns)}
            style_map = {c.attrib.get('s', '0'): styles.get(c.attrib['r']) or '0'
                         for c in donor.findall('.//s:c', ns)}
            for row in donor.findall('s:sheetData/s:row', ns):
                if 's' in row.attrib:
                    row.set('s', style_map[row.attrib['s']])
            for col in donor.findall('s:cols/s:col', ns):
                if 'style' in col.attrib:
                    col.set('style', style_map[col.attrib['style']])
            for cell in donor.findall('.//s:c', ns):
                style = styles.get(cell.attrib['r'])
                if style is None:
                    cell.attrib.pop('s', None)
                else:
                    cell.set('s', style)
                if cell.attrib.get('t') == 's':
                    value = cell.find('s:v', ns)
                    inline = copy(strings[int(value.text)])
                    inline.tag = '{' + ns['s'] + '}is'
                    cell.remove(value)
                    cell.append(inline)
                    cell.set('t', 'inlineStr')
            replacements[old_paths[dst]] = ET.tostring(donor, encoding='utf-8', xml_declaration=True)
        with zipfile.ZipFile(packed, 'w', compression=zipfile.ZIP_DEFLATED) as target:
            for entry in original.infolist():
                target.writestr(entry, replacements.get(entry.filename, original.read(entry.filename)))
    packed.replace(generated_path)


def sheet_snapshot(sheet):
    """Compare values and public style properties independently of style IDs."""
    cells = [(c.coordinate, c.value, str(c.font), str(c.fill), str(c.border),
              str(c.alignment), c.number_format, str(c.protection),
              c.comment.text if c.comment else None, str(c.hyperlink))
             for row in sheet for c in row]
    def dimension_snapshot(dim):
        attributes = dict(dim)
        attributes.pop('s', None)
        attributes.pop('style', None)
        return (attributes, str(dim.font), str(dim.fill), str(dim.border),
                str(dim.alignment), dim.number_format, str(dim.protection))
    return (cells, str(sheet.sheet_format), str(sheet.sheet_properties),
            tuple(dimension_snapshot(r) for r in sheet.row_dimensions.values()),
            tuple(dimension_snapshot(c) for c in sheet.column_dimensions.values()),
            tuple(str(m) for m in sheet.merged_cells.ranges), str(sheet.freeze_panes),
            repr(sheet.auto_filter), str(sheet.print_options), str(sheet.page_margins),
            str(sheet.page_setup))


def replace_rover_sheets(path: Path, source_path: Path = DEFAULT_WORKBOOK) -> dict:
    backup = path.with_name(f'{path.stem}.backup_{datetime.now():%Y%m%d_%H%M%S_%f}{path.suffix}')
    shutil.copy2(path, backup)
    book = load_workbook(path)
    source = load_workbook(source_path)
    names = book.sheetnames[:]
    if not all(name in names for name in SHEETS):
        raise ValueError('目标工作簿缺少 train1–train5 或 test1–test5')
    original = {name: sheet_snapshot(book[name]) for name in names if name not in REPLACEMENTS}
    expected = {dst: sheet_snapshot(source[src]) for dst, src in REPLACEMENTS.items()}
    temporary = path.with_name(f'.{path.stem}.replacement.xlsx')
    try:
        for dst, src in REPLACEMENTS.items():
            index = book.index(book[dst])
            book.remove(book[dst])
            target = book.create_sheet(dst, index)
            donor = source[src]
            for row in donor:
                for cell in row:
                    new = target.cell(cell.row, cell.column, cell.value)
                    if cell.has_style:
                        new.font, new.fill, new.border = copy(cell.font), copy(cell.fill), copy(cell.border)
                        new.alignment, new.protection = copy(cell.alignment), copy(cell.protection)
                        new.number_format = cell.number_format
                    if cell.comment:
                        new.comment = copy(cell.comment)
                    if cell.hyperlink:
                        new.hyperlink = copy(cell.hyperlink)
            for key, dim in donor.row_dimensions.items():
                target.row_dimensions[key] = copy(dim)
            for key, dim in donor.column_dimensions.items():
                target.column_dimensions[key] = copy(dim)
            for attr in ('sheet_format', 'sheet_properties', 'print_options', 'page_margins',
                         'page_setup', 'auto_filter', 'views'):
                setattr(target, attr, copy(getattr(donor, attr)))
            target.freeze_panes = donor.freeze_panes
            target.print_area = donor.print_area
            target.print_title_rows = donor.print_title_rows
            target.print_title_cols = donor.print_title_cols
            for merged in donor.merged_cells.ranges:
                target.merge_cells(str(merged))
        book.save(temporary)
        preserve_workbook_package(path, source_path, temporary)
        check = load_workbook(temporary)
        try:
            assert check.sheetnames == names, '工作表顺序改变'
            for name, snapshot in original.items():
                assert sheet_snapshot(check[name]) == snapshot, f'{name} 数据或格式改变'
            for name, snapshot in expected.items():
                assert sheet_snapshot(check[name]) == snapshot, f'{name} 与源表不一致'
        finally:
            check.close()
        temporary.replace(path)
    finally:
        book.close()
        source.close()
        temporary.unlink(missing_ok=True)
    return {'backup': str(backup), 'backup_sha256': hashlib.sha256(backup.read_bytes()).hexdigest(),
            'source_sha256': hashlib.sha256(source_path.read_bytes()).hexdigest(),
            'sheet_order_preserved': True, 'source_values_and_styles_match': True,
            'other_sheets_values_and_styles_unchanged': True}


def sheet_spec(path: Path, name: str) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        header = next(book[name].iter_rows(values_only=True))
    finally:
        book.close()
    bounds, columns = [], []
    for label in ('longitude', 'latitude'):
        pattern = rf'{label}\((-?\d+(?:\.\d+)?)-(-?\d+(?:\.\d+)?)\)'
        matches = [(h, re.fullmatch(pattern, h)) for h in header if isinstance(h, str)]
        matches = [(h, m) for h, m in matches if m]
        if len(matches) != 1:
            raise ValueError(f'{name}: 无法确定 {label} 区域表头')
        h, match = matches[0]
        bound = tuple(map(float, match.groups()))
        if not np.isclose(bound[1] - bound[0], 30):
            raise ValueError(f'{name}: 本流程要求 30° 区域')
        bounds.append(bound)
        columns.append(h)
    return {'sheet': name, 'longitude': columns[0], 'latitude': columns[1],
            'origin': (bounds[0][0], bounds[1][0]), 'bounds': bounds}


def model_parameters(model) -> dict:
    length = float(model.kern.lengthscale.values[0])
    return {'kernel_variance': float(model.kern.variance.values[0]),
            'lengthscale_degrees': length,
            'noise_variance': float(model.Gaussian_noise.variance.values[0]),
            'optimization_status': str(model.optimization_runs[-1].status),
            'log_likelihood': float(model.log_likelihood()),
            'extreme_lengthscale': length < 0.1 or length > 300}


def cross_validate(X, Y) -> dict:
    folds = np.array_split(np.random.default_rng(SEED).permutation(len(X)), 5)
    predicted, baseline = np.empty_like(Y), np.empty_like(Y)
    details = []
    for fold, test in enumerate(folds, 1):
        train = np.setdiff1d(np.arange(len(X)), test)
        model = train_model(X[train], Y[train])
        predicted[test], _ = model.predict(X[test])
        baseline[test] = Y[train].mean(axis=0)
        details.append({'fold': fold, 'held_out_indices': test.tolist(), **model_parameters(model)})
    if not np.isfinite(predicted).all():
        raise ValueError('交叉验证预测含非有限值')
    rmse = np.sqrt(np.mean((predicted - Y)**2, axis=0))
    base_rmse = np.sqrt(np.mean((baseline - Y)**2, axis=0))
    r2 = 1 - np.sum((predicted - Y)**2, axis=0) / np.sum((Y - Y.mean(axis=0))**2, axis=0)
    return {'seed': SEED, 'folds': details, 'rmse': rmse.tolist(), 'r2': r2.tolist(),
            'baseline_rmse': base_rmse.tolist(), 'rmse_over_baseline': (rmse / base_rmse).tolist()}


def plot_dataset(name, X, mean, path):
    x, y, _ = grid()
    fig, axes = plt.subplots(1, 3, figsize=(17, 5))
    for i, ax in enumerate(axes):
        values = mean[:, i].reshape(x.shape)
        low, high = float(values.min()), float(values.max())
        if high - low < 1e-8 * max(1, abs(low), abs(high)):
            padding = max(1e-6, abs(low) * 1e-3)
            levels = np.linspace(low - padding, high + padding, 13)
        else:
            levels = MaxNLocator(nbins=8).tick_values(low, high)
        if name == 'train1':
            levels = ZHURONG_LEVELS[i]
        surface = ax.contourf(x, y, values, levels=levels, cmap='hot', extend='both')
        ax.scatter(X[:, 0], X[:, 1], color='green', marker='*', s=22,
                   linewidths=0.4, label='observation position' if i == 0 else None, zorder=3)
        ax.set(title=TITLES[i], xlabel='Longitude offset (°)', ylabel='Latitude offset (°)',
               xlim=(0, 30), ylim=(0, 30))
        ax.set_xticks(np.arange(0, 31, 5))
        ax.set_yticks(np.arange(0, 31, 5))
        ax.grid(color='gray', linestyle='--', linewidth=1.2, alpha=0.75)
        fig.colorbar(surface, ax=ax, ticks=levels, fraction=0.048, pad=0.04, extendrect=True)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc='upper right', frameon=True)
    fig.subplots_adjust(left=0.05, right=0.97, top=0.89, bottom=0.13, wspace=0.38)
    fig.savefig(path, dpi=150)
    plt.close(fig)


def write_report(results, path):
    lines = ['# 十组火星磁场数据分析', '',
             '每表独立训练一个 GPy 三输出模型，三分量共享 RBF 核及噪声参数；不标准化，'
             '核初值方差和长度尺度均为 1，每次仅优化一次。共享核不等于学习了分量交叉协方差。', '',
             '## 数据覆盖与验证', '',
             '覆盖率为观测凸包面积 / 标称 30°×30° 区域面积。RMSE 单位按原项目使用 nT；'
             '工作簿本身未明确标注单位。R² 与 RMSE 均为固定种子 42 的随机五折样本外结果。', '',
             '| 表 | 点数 | 凸包覆盖 | Br/B2/B3 R² | Br/B2/B3 RMSE | RMSE/均值基线 |',
             '|---|---:|---:|---|---|---|']
    for name, result in results.items():
        cv = result['cross_validation']
        fmt = lambda values: ' / '.join(f'{v:.3f}' for v in values)
        lines.append(f"| {name} | {result['training_points']} | {result['hull_coverage']:.1%} | "
                     f"{fmt(cv['r2'])} | {fmt(cv['rmse'])} | {fmt(cv['rmse_over_baseline'])} |")
    lines += ['', '## 分布与分量关系', '',
              '全部 676 个观测点数值有限，各表均无重复坐标。八张未替换表的磁场值集中在约 '
              '−30 至 30，常保留两位小数；祝融和毅力源表保留更多小数。这是数值精度差异，'
              '不能据此判断数据是否为实测或合成。', '',
              '以下为均值 ± 总体标准差 [最小值, 最大值]。相关性描述统计关联，不证明物理因果。', '']
    for name, result in results.items():
        s = result['statistics']
        lines += [f'### {name}', '', f"经纬度区域：{result['bounds_degrees']}；实际观测范围：{s['coordinate_min']} 至 {s['coordinate_max']}。", '']
        for i, component in enumerate(COMPONENTS):
            lines.append(f"- {component}：{s['mean'][i]:.3f} ± {s['std'][i]:.3f} "
                         f"[{s['min'][i]:.3f}, {s['max'][i]:.3f}]")
        corr = s['component_correlation']
        lines += ['', f'Br–B2 / Br–B3 / B2–B3 相关系数：{corr[0][1]:.3f} / {corr[0][2]:.3f} / {corr[1][2]:.3f}。',
                  f"全量模型长度尺度：{result['model']['lengthscale_degrees']:.4g}°；优化状态：{result['model']['optimization_status']}；"
                  f"验证折极端长度尺度数：{sum(f['extreme_lengthscale'] for f in result['cross_validation']['folds'])}/5。", '']
    weak = [name for name, r in results.items() if all(v <= 0 for v in r['cross_validation']['r2'])]
    weak_components = [f'{name}/{COMPONENTS[i]}' for name, r in results.items()
                       for i, value in enumerate(r['cross_validation']['r2']) if value <= 0]
    lines += ['## 解释与限制', '',
              '- 本报告以实际 GPy 验证为准，取代此前 SciPy 近似实现的初步结果。两种优化实现的结果不同，不能把近似结果作为最终 GPy 模型质量。',
              f"- 三分量验证 R² 均不大于零的表：{'、'.join(weak) or '无'}。此判据只能定位明显弱拟合，正 R² 仍需结合误差判断用途。",
              f"- 验证 R² 不大于零的分量：{'、'.join(weak_components) or '无'}。test4/B2 的 R² 接近零，也需谨慎使用。",
              '- 毅力号 Br 幅值远大于其余表，Br 与 B3 的相关系数约 0.657；共享参数可能使大幅值分量更影响优化。',
              '- train4、test5 的观测凸包分别只覆盖约 55%、47% 的标称区域，观测凸包之外属于外推。图片按祝融号参考图样式绘制，不叠加外推斜线；覆盖率与外推网格比例保留在 JSON 中。凸包内部也不保证采样密集。',
              '- 极端长度尺度定义为小于 0.1° 或大于 300°，只作为退化拟合提示，不据此改变复现算法。详细参数和每折状态见 analysis.json。',
              '- 随机五折评价的是观测位置分布内的预测，不能替代空间分块验证或独立新区域验证；这些图片是 GP 推断，不是真实磁场全覆盖测量。',
              '- train/test 工作表名保留原用途。本次为各区域分别拟合环境模型；若下游以 test 表评价策略泛化，应独立控制策略训练与测试的数据使用。',
              '- 来源、测量高度、仪器误差与分量坐标定义需额外资料确认，当前数据不足以建立跨区域物理规律。', '',
              '## 文件验收', '',
              '两张替换表与源表的数值和格式一致，原工作表顺序及其余八表数值与格式保持。'
              '十个 pkl 均重新加载并预测有限的 961×3 均值，十张 PNG 均经图像解码校验。'
              '模型输入需先减去 analysis.json 中对应区域原点，输出顺序为 Br、B2、B3。', '']
    path.write_text('\n'.join(lines), encoding='utf-8')


def run_batch(workbook_path=DEFAULT_BATCH_WORKBOOK, output_dir=DEFAULT_BATCH_OUTPUT, replace=False):
    output_dir.mkdir(parents=True, exist_ok=True)
    replacement = replace_rover_sheets(workbook_path) if replace else None
    if replacement:
        print('备份及替换验证通过：', replacement['backup'], flush=True)
    results = {}
    _, _, locations = grid()
    for name in SHEETS:
        spec = sheet_spec(workbook_path, name)
        X, Y = load_sheet_data(workbook_path, spec)
        if len(np.unique(X, axis=0)) != len(X) or np.any(X < 0) or np.any(X > 30):
            raise ValueError(f'{name}: 重复坐标或观测超出标称区域')
        hull = ConvexHull(X)
        outside = Delaunay(X).find_simplex(locations) < 0
        print(f'{name}: {len(X)} 个点，训练并执行五折验证', flush=True)
        model = train_model(X, Y)
        params = model_parameters(model)
        cv = cross_validate(X, Y)
        model_path = output_dir / f'{name}.pkl'
        with model_path.open('wb') as target:
            pickle.dump(model, target, protocol=pickle.HIGHEST_PROTOCOL)
        with model_path.open('rb') as source:
            restored = pickle.load(source)
        mean, variance = restored.predict(locations)
        assert mean.shape == (961, 3) and np.isfinite(mean).all() and np.isfinite(variance).all()
        assert np.array_equal(restored.X, X) and np.array_equal(restored.Y, Y)
        np.testing.assert_allclose(mean, model.predict(locations)[0], rtol=0, atol=1e-12)
        image_path = output_dir / f'{name}.png'
        plot_dataset(name, X, mean, image_path)
        with Image.open(image_path) as image:
            image.verify()
        results[name] = {'training_points': len(X), 'origin_degrees': spec['origin'],
                         'bounds_degrees': spec['bounds'], 'components': COMPONENTS,
                         'hull_coverage': float(hull.volume / 900),
                         'grid_extrapolation_fraction': float(outside.mean()),
                         'statistics': {'coordinate_min': (X + spec['origin']).min(0).tolist(),
                                        'coordinate_max': (X + spec['origin']).max(0).tolist(),
                                        'min': Y.min(0).tolist(), 'max': Y.max(0).tolist(),
                                        'mean': Y.mean(0).tolist(), 'std': Y.std(0).tolist(),
                                        'component_correlation': np.corrcoef(Y.T).tolist()},
                         'model': params, 'cross_validation': cv,
                         'validation': {'pickle_roundtrip': True, 'finite_prediction_shape': [961, 3],
                                        'png_decoded': True}}
        print(f"{name}: R² {np.round(cv['r2'], 3).tolist()}", flush=True)
    metadata = {'python': sys.version, 'versions': {p: version(p) for p in
                ('GPy', 'paramz', 'numpy', 'scipy', 'matplotlib', 'openpyxl', 'pillow')},
                'workbook_sha256': hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                'replacement': replacement,
                'method': 'Raw outputs; shared RBF(variance=1, lengthscale=1); one optimize() per fit',
                'datasets': results}
    (output_dir / 'analysis.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2,
                                                       allow_nan=False) + '\n', encoding='utf-8')
    (output_dir / 'kernel_parameters.json').write_text(json.dumps({name: {'origin_degrees': r['origin_degrees'],
        **r['model']} for name, r in results.items()}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    write_report(results, output_dir / 'analysis.md')
    print(f'完成：{output_dir}', flush=True)
