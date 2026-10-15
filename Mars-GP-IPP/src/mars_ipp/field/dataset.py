"""Original workbook columns and row order; coordinates are local degrees."""
import numpy as np
from openpyxl import load_workbook

COMPONENTS = ('Br', 'B2', 'B3')
SCENARIOS = {
    'train1': {'sheet': 'Zhurong_Magnetic_Field', 'columns': ('longitude(95-125)', 'latitude(10-40)'), 'origin': (95., 10.)},
    'zhurong': {'sheet': 'Zhurong_Magnetic_Field', 'columns': ('longitude(95-125)', 'latitude(10-40)'), 'origin': (95., 10.)},
    'perseverance': {'sheet': 'Perseverance_Magnetic_Field', 'columns': ('longitude(65-95)', 'latitude(5-35)'), 'origin': (65., 5.)},
}


def load_dataset(path, scenario):
    spec = SCENARIOS[scenario]
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book[spec['sheet']].iter_rows(values_only=True)
        header = next(rows)
        indices = [header.index(c) for c in (*spec['columns'], *COMPONENTS)]
        data = []
        for row in rows:
            selected = [row[i] for i in indices]
            if all(v is None for v in selected):
                continue
            data.append([float(v) for v in selected])
        data = np.asarray(data, dtype=float)
        if data.ndim != 2 or data.shape[1] != 5 or not np.isfinite(data).all():
            raise ValueError('Invalid workbook data')
        return data[:, :2] - spec['origin'], data[:, 2:]
    finally:
        book.close()


def evaluation_grid(size=31, extent=30.):
    x, y = np.meshgrid(np.linspace(0, extent, size), np.linspace(0, extent, size), indexing='xy')
    return np.column_stack((x.ravel(order='C'), y.ravel(order='C')))
